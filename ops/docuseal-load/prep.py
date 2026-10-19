#!/usr/bin/env python3
"""Rev D package -> filled DOCX, PDF parts, and a DocuSeal template payload per part.

Fills what facts.local.json states. Every other blank becomes a line wide enough for its answer,
owned by the party who knows it (Keyamo's state -> Keyamo, a nominee's address -> that nominee,
money terms -> the sender). Converts with LibreOffice, splits exhibits that are signed by other
people or in ink into their own parts, then finds every signature and fill line in each part
(underscore runs and rules drawn as borders) and places a DocuSeal field on it. Run with python3 -I.

  python3 -I prep.py PACKAGE.zip|DIR|FILE.docx --facts facts.local.json --out WORK
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import numpy as np
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw
from pypdf import PdfReader, PdfWriter

PT_GAP = 30          # word gap (pt) that splits a pdftotext line into column segments
SIG_H, TXT_H = 26, 14
LABELS = {"by", "signature", "sign", "/s/", "name", "printed name", "print name", "title", "date", "initials",
          "its", "business address", "address", "acknowledged by the employee"}
BLANK_RE = re.compile(r"_{4,}|\[_+\]")
REPORT_RE = re.compile(r"\[[^\]]{1,80}\]|_{4,}")
NS = "{http://www.w3.org/1999/xhtml}"


# ---------- fill ----------

def iter_paragraphs(doc):
    """Yield (paragraph, caption, row label) for body, tables (recursively) and headers/footers.
    caption = first non-empty paragraph of the enclosing cell, else the last all-caps body line;
    row label = text of the first cell in the row (what a value cell answers)."""
    def cell_pars(cell, rowlabel, width):
        cap = next((p.text.strip() for p in cell.paragraphs if p.text.strip()), "")
        for p in cell.paragraphs:
            yield p, cap, rowlabel, width
        for t in cell.tables:
            yield from table_pars(t)

    def table_pars(t):
        for row in t.rows:
            seen, label = set(), ""
            for i, cell in enumerate(row.cells):
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                if i == 0:
                    label = cell.text.strip()
                width = cell.width.pt - 12 if cell.width else None
                yield from cell_pars(cell, label if i else "", width)

    cap = ""
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(child, doc)
            if p.text.strip() and is_caption(p.text.strip()):
                cap = p.text.strip()
            yield p, cap, "", None
        elif tag == "tbl":
            yield from table_pars(Table(child, doc))
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            for p in hf.paragraphs:
                yield p, "", "", None


def para_replace(p, old, new, limit=None):
    n, start = 0, 0
    while limit is None or n < limit:
        runs = p.runs
        full = "".join(r.text for r in runs)
        i = full.find(old, start)
        if i < 0 or not runs:
            return n
        pos, sr = 0, None
        for k, r in enumerate(runs):
            L = len(r.text)
            if sr is None and i < pos + L:
                sr, so = k, i - pos
            if sr is not None and i + len(old) <= pos + L:
                er, eo = k, i + len(old) - pos
                break
            pos += L
        else:
            return n  # match spans non-run content (hyperlink/field); leave it
        if sr == er:
            runs[sr].text = runs[sr].text[:so] + new + runs[sr].text[eo:]
        else:
            runs[sr].text = runs[sr].text[:so] + new
            for k in range(sr + 1, er):
                runs[k].text = ""
            runs[er].text = runs[er].text[eo:]
        start, n = i + len(new), n + 1
    return n


def applies(rule, key, cap):
    if rule.get("docs") and not any(key.startswith(d) for d in rule["docs"]):
        return False
    return not rule.get("caption") or re.search(rule["caption"], cap, re.I)


def fill_docx(src, dst, key, facts):
    """fill: exact-text replacements of what is known. blanks: per [___] occurrence, the first rule whose
    ctx regex matches the text just before it (row label included) sets a value, or a line width."""
    doc = Document(src)
    applied, opened = [], []
    default_w = facts.get("blank_width", 14)
    def fit(w, p, i, avail):
        """Cap a line so it fits on one line of its table cell (underscore ~5.5pt, text ~5pt/char)."""
        if not avail:
            return w
        used = (len(p.text[:i].split("\n")[-1]) * 5.0) % avail
        room = int((avail - used) / 5.5)
        return max(4, min(w, room if room >= 6 else int(avail / 5.5)))

    for p, cap, rowlabel, avail in iter_paragraphs(doc):
        for r in facts.get("fill", []):
            if applies(r, key, cap) and r["find"] in p.text:
                new = r["replace"] if "replace" in r else "_" * fit(r["blank"], p, p.text.find(r["find"]), avail)
                n = para_replace(p, r["find"], new)
                if n:
                    (applied if "replace" in r else opened).append(
                        {"find": r["find"], "to": new if "replace" in r else f"{r['blank']}-char line", "caption": cap, "count": n})
        while "[___]" in p.text:
            i = p.text.find("[___]")
            ctx = re.sub(r"\s+", " ", (rowlabel + " | " if rowlabel else "") + p.text[max(0, i - 120):i])
            rule = next((r for r in facts.get("blanks", []) if not r.get("x") and applies(r, key, cap) and re.search(r["ctx"], ctx, re.I)), None)
            if rule and "value" in rule:
                para_replace(p, "[___]", rule["value"], limit=1)
                applied.append({"find": "[___]", "to": rule["value"], "context": ctx[-90:], "caption": cap})
            else:
                w = fit(rule.get("width", default_w) if rule else default_w, p, i, avail)
                if not para_replace(p, "[___]", "_" * w, limit=1):
                    break
                opened.append({"context": ctx[-90:], "caption": cap, "width": w})
    leftovers = [{"placeholder": m.group(0)[:160], "caption": cap[:60], "context": p.text.strip()[:160]}
                 for p, cap, _, _ in iter_paragraphs(doc) for m in re.finditer(r"\[[^\]_ X][^\]]{0,400}\]", p.text)]
    doc.save(dst)
    return applied, opened, leftovers


# ---------- detect ----------

def is_caption(t):
    letters = [c for c in t if c.isalpha()]
    if len(letters) < 3 or "_" in t or t.endswith(":"):
        return False
    if t.split(":")[0].strip().lower() in LABELS:
        return False
    return sum(c.isupper() for c in letters) / len(letters) > 0.8


def pdf_segments(pdf, tmp):
    html = tmp / (pdf.stem + ".bbox.html")
    subprocess.run(["pdftotext", "-bbox-layout", str(pdf), str(html)], check=True)
    root = ET.parse(html).getroot()
    pages = []
    for pg in root.iter(NS + "page"):
        W, H = float(pg.get("width")), float(pg.get("height"))
        segs = []
        for line in pg.iter(NS + "line"):
            words = [dict(t=w.text or "", x0=float(w.get("xMin")), y0=float(w.get("yMin")),
                          x1=float(w.get("xMax")), y1=float(w.get("yMax"))) for w in line.iter(NS + "word")]
            cur = []
            for w in words:
                if cur and w["x0"] - cur[-1]["x1"] > PT_GAP:
                    segs.append(cur)
                    cur = []
                cur.append(w)
            if cur:
                segs.append(cur)
        segs = [dict(words=s, text=" ".join(w["t"] for w in s), x0=s[0]["x0"], x1=s[-1]["x1"],
                     y0=min(w["y0"] for w in s), y1=max(w["y1"] for w in s)) for s in segs]
        pages.append(dict(W=W, H=H, segs=segs))
    return pages


def drawn_rules(pdf, pno, W, H, tmp, dpi=100, min_in=1.2):
    """Solid black horizontal rules (signature lines drawn as paragraph borders). Colored rules
    (headers, headings, table grids) are excluded by the near-black test."""
    stem = tmp / f"{pdf.stem}.r{pno}"
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-f", str(pno), "-l", str(pno), "-singlefile", str(pdf), str(stem)], check=True)
    a = np.asarray(Image.open(f"{stem}.png").convert("RGB")).astype(int)
    dark = (a.max(axis=2) < 50)
    out = []
    for y in range(dark.shape[0]):
        row = np.concatenate(([0], dark[y].astype(int), [0]))
        d = np.diff(row)
        for s, e in zip(np.where(d == 1)[0], np.where(d == -1)[0]):
            if e - s >= min_in * dpi and dark[max(y - 3, 0):y + 4, s:e].mean(axis=1).max() > 0.9 \
                    and (y < 4 or dark[y - 4, s:e].mean() < 0.3) and (y + 4 >= dark.shape[0] or dark[y + 4, s:e].mean() < 0.3):
                out.append((s * W / a.shape[1], e * W / a.shape[1], y * H / a.shape[0]))
    merged = []
    for x0, x1, y in sorted(out, key=lambda t: (round(t[0]), t[2])):
        if merged and abs(merged[-1][0] - x0) < 3 and abs(y - merged[-1][2]) < 2.5:
            continue
        merged.append((x0, x1, y))
    return merged


def classify(label):
    l = label.lower()
    if "initial" in l:
        return "initials"
    if "date" in l:
        return "date"
    if l.startswith("by") or "/s/" in l or "signature" in l or l.startswith("sign") or "acknowledged" in l:
        return "signature"
    return "text"


def block_context(segs, seg, cap_window=140, above=90, below=110):
    """Party evidence around a line, same column: the all-caps caption heading its block, the lines
    nearby ordered by distance, and the printed name of the block's signer."""
    col = seg["x0"]
    same = [s for s in segs if s is not seg and abs(s["x0"] - col) < 45]
    caps = [s for s in same if s["y1"] <= seg["y0"] + 1 and seg["y0"] - s["y1"] < cap_window and is_caption(s["text"])]
    cap = max(caps, key=lambda s: s["y1"])["text"] if caps else ""
    near = [(seg["y0"] - s["y1"], s["text"]) for s in same if 0 <= seg["y0"] - s["y1"] < above]
    near += [(s["y0"] - seg["y1"], s["text"]) for s in same if 0 <= s["y0"] - seg["y1"] < below]
    near = [t for _, t in sorted(near)]
    name = next((m.group(2).strip() for t in near for m in [re.match(r"^(Name|By):\s*(.+)$", t)]
                 if m and not BLANK_RE.search(m.group(2)) and not m.group(2).startswith("[")), "")
    name = re.sub(r"\s*\[.*$", "", name)
    return cap, near, name


def find_role(cap, near, name, roles, key, label=""):
    """roles: [pattern, role] or [pattern, role, [doc keys]]. Role may use \\1 groups and {name}.
    A block caption that names a party wins; otherwise the nearest matching line."""
    live = [r for r in roles if len(r) < 3 or not r[2] or key in r[2]]
    for text in [t for t in [cap, label] if t] + near:
        for r in live:
            m = re.search(r[0], text, re.I)
            if m:
                role = m.expand(r[1])
                return role.replace("{name}", name) if "{name}" not in role or name else None
    return None


def blank_owner(ctx, rules, key, x=None, page=None):
    ctx = re.sub(r"\s+", " ", ctx)
    for r in rules:
        if r.get("x") and (x is None or not r["x"][0] <= x <= r["x"][1]):
            continue
        if r.get("page") and page != r["page"]:
            continue
        if (not r.get("docs") or any(key.startswith(d) for d in r["docs"])) and (r.get("owner") or r.get("name")) \
                and re.search(r["ctx"], ctx, re.I):
            return r.get("owner"), r.get("name"), r
    return None, None, {}


def detect_fields(pdf, pages, facts, key, tmp, skip_pages=()):
    roles = facts["roles"]
    sender = facts.get("sender_by_doc", {}).get(key, facts["sender_role"])
    fields, open_items = [], []
    for pno, pg in enumerate(pages, 1):
        if pno in skip_pages:
            continue
        W, H, segs = pg["W"], pg["H"], pg["segs"]
        hits = []
        for seg in segs:
            ws = seg["words"]
            for i, w in enumerate(ws):
                m = BLANK_RE.search(w["t"])
                if not m:
                    continue
                cw = (w["x1"] - w["x0"]) / max(len(w["t"]), 1)
                x0 = w["x0"] + cw * m.start()
                x1 = w["x0"] + cw * (len(w["t"].rstrip(".,;)")))
                label = (" ".join(v["t"] for v in ws[:i]) + " " + w["t"][:m.start()]).strip()
                hits.append((seg, x0, x1, w["y1"], label, "u"))
        words = [w for s in segs for w in s["words"]]
        for x0, x1, y in drawn_rules(pdf, pno, W, H, tmp):
            if any(w["x0"] < x1 and x0 < w["x1"] and y - 18 < w["y1"] <= y + 8 for w in words):
                continue  # underscores, or cell text right above: not a drawn signature line
            seg = dict(words=[], text="", x0=x0, x1=x1, y0=y - 1, y1=y)
            hits.append((seg, x0, x1, y + 1, "", "r"))
        for seg in segs:
            ws = seg["words"]
            for i, w in enumerate(ws):
                box = None
                if w["t"] == "[]":
                    box = (w["x0"], w["x1"])
                elif w["t"] == "[" and i + 1 < len(ws) and ws[i + 1]["t"].startswith("]") and ws[i + 1]["x0"] - w["x1"] < 14:
                    box = (w["x0"], ws[i + 1]["x1"])
                if not box:
                    continue
                after = re.sub(r"_{2,}", "", " ".join(v["t"] for v in ws[i + 1:i + 6]).lstrip("] ").split("[")[0]).strip(" :")
                row_left = " ".join(s["text"] for s in segs if s["x1"] < box[0] and s["y0"] - 12 < seg["y1"] and seg["y0"] < s["y1"] + 12)
                ctx = re.sub(r"\s+", " ", re.sub(r"_{2,}", " ", f"{row_left} {' '.join(v['t'] for v in ws[:i])}"))
                role, fname, _r = blank_owner(ctx + " [ ] " + after, facts.get("checks", []), key, (box[0] + box[1]) / 2 / W, pno)
                side = min(box[1] - box[0], 9)
                area = dict(x=round(box[0] / W, 5), y=round((w["y1"] - side - 1) / H, 5),
                            w=round(side / W, 5), h=round(side / H, 5), page=pno)
                fields.append(dict(type="checkbox", role=role or sender, label=after, caption="", area=area,
                                   name=fname or ("Check: " + (after or row_left)[:40]), same=False, context=(ctx + " [ ] " + after)[-140:],
                                   note=None, in_block=False))
        for seg, x0, x1, ybase, label, kind in hits:
            lkey = label.split(":")[0].strip().lower() if ":" in label else ""
            last = (label.split(":")[-2].split() or [""])[-1].lower() if label.count(":") else ""
            left = [s for s in segs if s is not seg and s["x1"] < x0 and s["y0"] - 12 < seg["y1"] and seg["y0"] < s["y1"] + 12]
            mid = lambda s: (s["y0"] + s["y1"]) / 2
            dist = [abs(mid(s) - mid(seg)) for s in left]
            row_left = " ".join(s["text"] for s, d in sorted(zip(left, dist), key=lambda t: (t[0]["y0"], t[0]["x0"]))
                                if d <= min(dist) + 3) if left else ""
            beside = bool(left)
            below = [s for s in segs if 0 <= s["y0"] - seg["y1"] < 18 and abs(s["x0"] - seg["x0"]) < 20]
            named = bool(below) and below[0]["text"][:1].isupper() and "_" not in below[0]["text"] and not is_caption(below[0]["text"])
            bare = label == "" and (kind == "r" or (not beside and named and x1 - x0 >= 108 and not seg["text"].strip("_ ")))
            in_block = bare or lkey in LABELS or last in LABELS or label.lower().startswith("/s/")
            ftype = "signature" if bare else classify(label) if in_block else "text"
            cap, near, name = block_context(segs, seg)
            role = find_role(cap, near, name, roles, key, label) if in_block else None
            prev = [s["text"] for s in segs if 0 < seg["y0"] - s["y1"] < 16 and abs(s["x0"] - seg["x0"]) < 260]
            ctx = " ".join(prev[-1:] + [row_left, label])
            fname, rule = None, {}
            if role is None:
                role, fname, rule = blank_owner(ctx, facts.get("blanks", []), key, (x0 + x1) / 2 / W, pno)
                if role is None and fname:
                    role = sender
            ftype = rule.get("type", ftype)
            same = rule.get("same", False)
            note = None
            if role is None:
                role = sender
                note = None if not in_block else f"no party match (caption '{cap[:40]}')"
            h = SIG_H if ftype in ("signature", "initials") else TXT_H
            area = dict(x=round(x0 / W, 5), y=round(max(ybase - 1 - h, 0) / H, 5),
                        w=round((x1 - x0) / W, 5), h=round(h / H, 5), page=pno)
            if not in_block and not fname:
                fname = f"p{pno} " + (" ".join(re.sub(r"_{2,}|\[ \]", " ", ctx).split()[-5:]) or "blank")
            fields.append(dict(type=ftype, role=role, label=label, caption=cap, area=area, name=fname, same=same,
                               context=(ctx or seg["text"])[-140:], note=note, in_block=in_block))
            if note:
                open_items.append(dict(page=pno, label=label or ("(drawn line)" if kind == "r" else "(blank)"), context=ctx[-140:], note=note))
    return fields, open_items


def field_name(f):
    if f["name"]:
        return f"{f['role']}: {f['name']}"
    if not f["in_block"]:
        words = re.sub(r"\[_+\]|_{4,}", " ", f["context"]).split()[-5:]
        return f"{f['role']}: " + (" ".join(words) or "blank")
    base = {"signature": "Signature", "date": "Date", "initials": "Initials"}.get(f["type"]) \
        or ((f["label"].split(":")[-2].split() or [""])[-1].title() if f["label"].count(":") else f["label"].title()) or "Text"
    return f"{f['role']} {base}"


def to_payload(name, ext_id, folder, pdf_path, fields, role_order, date_format):
    order = {r: i for i, r in enumerate(role_order)}
    fields = sorted(fields, key=lambda f: (order.get(f["role"], 99), f["area"]["page"], f["area"]["y"], f["area"]["x"]))
    seen, out, shared = {}, [], {}
    for f in fields:
        nm = field_name(f)
        if f.get("same") and (f["role"], nm) in shared:
            shared[(f["role"], nm)]["areas"].append(f["area"])  # one value, shown wherever it recurs
            continue
        seen[nm] = seen.get(nm, 0) + 1
        if seen[nm] > 1:
            nm = f"{nm} {seen[nm]}"
        fd = {"name": nm, "type": f["type"], "role": f["role"], "required": f["type"] != "checkbox", "areas": [f["area"]]}
        if f["type"] == "date":
            fd["preferences"] = {"format": date_format}
        if not f["in_block"]:
            fd["description"] = f"Fill in: {f['context'][-120:]}"
        if f.get("same"):
            shared[(f["role"], nm)] = fd
        out.append(fd)
    return {"name": name, "external_id": ext_id, "folder_name": folder,
            "documents": [{"name": name, "file": f"@file:{pdf_path}", "fields": out}]}


# ---------- parts ----------

def split_parts(pdf, key, facts, outdir):
    """[{suffix, name, pdf, esign, note}] for this document; one part (the whole PDF) when not split.
    Parts are cut on headings, not page numbers (filled lines move page breaks): a page whose top
    matches one of a part's "starts" regexes begins that part, and later pages stay with it until
    another part starts. The first part also holds every page before the first match."""
    specs = facts.get("parts", {}).get(key)
    if not specs:
        return [dict(suffix="", name="", pdf=pdf, esign=True, note="")]
    r = PdfReader(pdf)
    owner, cur = [], 0
    for i in range(len(r.pages)):
        top = subprocess.run(["pdftotext", "-f", str(i + 1), "-l", str(i + 1), str(pdf), "-"],
                             capture_output=True, text=True).stdout[:400]
        for k, s in enumerate(specs):
            if any(re.search(x, top) for x in s.get("starts", [])):
                cur = k
        owner.append(cur)
    parts = []
    for k, s in enumerate(specs):
        w = PdfWriter()
        for i in [i for i, o in enumerate(owner) if o == k]:
            w.add_page(r.pages[i])
        if not w.pages:
            continue
        out = outdir / f"{pdf.stem}.{s['suffix']}.pdf"
        with open(out, "wb") as fh:
            w.write(fh)
        parts.append(dict(suffix=s["suffix"], name=s["name"], pdf=out, esign=s.get("esign", True), note=s.get("note", "")))
    return parts


# ---------- QA ----------

COLORS = ["#d1495b", "#2e86ab", "#3c9d5d", "#e09f3e", "#7b4fa0", "#1b998b", "#555555", "#c05299"]


def overlay(pdf, fields, role_order, outdir, dpi=70):
    roles = role_order + sorted({f["role"] for f in fields} - set(role_order))
    col = {r: COLORS[i % len(COLORS)] for i, r in enumerate(roles)}
    shots = []
    for p in sorted({f["area"]["page"] for f in fields}):
        stem = outdir / f"{pdf.stem}.p{p}"
        subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-f", str(p), "-l", str(p), "-singlefile", str(pdf), str(stem)], check=True)
        im = Image.open(f"{stem}.png").convert("RGB")
        d = ImageDraw.Draw(im)
        for f in fields:
            a = f["area"]
            if a["page"] != p:
                continue
            box = [a["x"] * im.width, a["y"] * im.height, (a["x"] + a["w"]) * im.width, (a["y"] + a["h"]) * im.height]
            d.rectangle(box, outline=col[f["role"]], width=2)
            d.text((box[0] + 2, box[1] - 10), f"{f['role']}:{f['type']}", fill=col[f["role"]])
        im.save(f"{stem}.png")
        shots.append(f"{stem}.png")
    return shots


def doc_key(stem):
    m = re.match(r"(UAT-\d+[A-Z]?)", stem)
    return m.group(1) if m else stem


def pretty_name(stem, rev):
    k = doc_key(stem)
    rest = re.sub(r"^UAT-\d+[A-Z]?_", "", stem)
    rest = re.sub(r"_?Rev[-_ ]?D[0-9.]*$", "", rest, flags=re.I).replace("_", " ")
    return f"{k} {rest} (Rev {rev})"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("package")
    ap.add_argument("--facts", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    facts = json.load(open(a.facts))
    out = Path(a.out)
    for d in ("src", "filled", "pdf", "parts", "qa", "tmp"):
        (out / d).mkdir(parents=True, exist_ok=True)

    src = Path(a.package)
    if src.suffix == ".zip":
        with zipfile.ZipFile(src) as z:
            for n in z.namelist():
                if n.endswith(".docx") and not n.startswith("__MACOSX"):
                    (out / "src" / Path(n).name).write_bytes(z.read(n))
        docs = sorted((out / "src").glob("*.docx"))
    elif src.suffix == ".docx":
        docs = [src]
    else:
        docs = sorted(src.rglob("*.docx"))

    rev = facts.get("rev", "D.2")
    role_order = facts.get("role_order", [])
    manifest, report = [], {"skipped": [], "ink": [], "docs": []}
    for d in docs:
        key = doc_key(d.stem)
        if any(key == s or d.stem.startswith(s) for s in facts.get("skip", [])):
            report["skipped"].append(d.name)
            continue
        filled = out / "filled" / d.name
        applied, opened, leftovers = fill_docx(d, filled, key, facts)
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(out / "pdf"), str(filled)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pdf = out / "pdf" / (d.stem + ".pdf")
        for part in split_parts(pdf, key, facts, out / "parts"):
            label = f"{key}{'-' + part['suffix'] if part['suffix'] else ''}"
            name = pretty_name(d.stem, rev) if not part["name"] else f"{key} {part['name']} (Rev {rev})"
            if not part["esign"]:
                report["ink"].append(dict(part=label, name=name, pdf=str(part["pdf"]), note=part["note"]))
                print(f"{label}: ink only -> {part['pdf'].name}")
                continue
            pages = pdf_segments(part["pdf"], out / "tmp")
            fields, open_items = detect_fields(part["pdf"], pages, facts, key, out / "tmp")
            if not fields:
                report["skipped"].append(f"{label} (no lines found)")
                continue
            ext = f"{label.lower()}-{rev.lower().replace('.', '')}"
            payload = to_payload(name, ext, facts.get("folder", "Rev D"), str(part["pdf"].resolve()), fields,
                                 role_order, facts.get("date_format", "MM/DD/YYYY"))
            pj = out / f"{label}.payload.json"
            json.dump(payload, open(pj, "w"), indent=1)
            shots = overlay(part["pdf"], fields, role_order, out / "qa")
            manifest.append(str(pj.resolve()))
            signers = sorted({f["role"] for f in fields if f["type"] == "signature"})
            report["docs"].append(dict(part=label, file=d.name, template=name, external_id=ext, pages=len(pages),
                                       note=part["note"], signers=signers,
                                       fields=[dict(name=x["name"], type=x["type"], role=x["role"], page=x["areas"][0]["page"])
                                               for x in payload["documents"][0]["fields"]],
                                       filled=applied, opened_lines=opened, placeholders_left=leftovers,
                                       needs_review=open_items, qa=shots))
            print(f"{label}: {len(fields)} fields, signers {signers}, {len(open_items)} to review")
    json.dump(manifest, open(out / "manifest.json", "w"), indent=1)
    json.dump(report, open(out / "report.json", "w"), indent=1)
    shutil.rmtree(out / "tmp", ignore_errors=True)
    print(f"{len(manifest)} templates ready, {len(report['ink'])} ink-only parts, {len(report['skipped'])} skipped")


if __name__ == "__main__":
    sys.exit(main())
